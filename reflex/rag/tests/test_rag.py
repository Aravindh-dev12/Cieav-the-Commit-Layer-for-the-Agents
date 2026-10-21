import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

H = {"x-rag-secret": "s3cret-value"}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("RAG_DB", str(tmp_path / "rag.db"))
    monkeypatch.setenv("RAG_SECRET", "s3cret-value")
    from app.main import app
    return TestClient(app)


def user(roles=(), clearance="internal"):
    return {"name": "t", "roles": list(roles), "clearance": clearance}


def put(c, **d):
    return c.post("/ingest/text", json=d, headers=H)


def hits(c, q, u):
    return c.post("/search", json={"question": q, "user": u}, headers=H).json()["hits"]


def tiny_pdf(text):
    stream = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET"
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream", "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out)); out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF".encode()


def test_fails_closed_without_secret(client):
    assert client.get("/health").status_code == 200
    assert client.post("/search", json={"question": "x", "user": user()}).status_code == 401


def test_versions_permissions_and_refusal(client):
    put(client, doc_key="REQ-PWR", body="Battery depth of discharge shall not exceed 60 percent.")
    assert put(client, doc_key="REQ-PWR", body="Battery depth of discharge shall not exceed 60 percent.").json()["unchanged"]
    assert put(client, doc_key="REQ-PWR", body="Battery depth of discharge shall not exceed 50 percent.").json()["version"] == 2
    put(client, doc_key="NOZ", level="restricted", roles=["export_cleared"], body="Propulsion nozzle thermal margin is 15 percent.")
    top = hits(client, "battery depth of discharge", user())[0]
    assert top["version"] == 2 and "50 percent" in top["text"]
    assert hits(client, "nozzle thermal margin", user()) == []  # invisible to a normal engineer
    assert hits(client, "nozzle thermal margin", user(["admin"], "restricted"))[0]["doc_key"] == "NOZ"
    assert hits(client, "who won the world cup", user()) == []  # unsupported question is refused
    assert client.post("/ingest/text", json={"doc_key": "X", "body": "y", "level": "secret"}, headers=H).status_code == 422


def test_docx_tables_and_pdf_are_parsed(client):
    d = Document()
    d.add_paragraph("Thermal control requirements.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "REQ-TH-4", "Radiator shall reject 400 W"
    buf = io.BytesIO(); d.save(buf)
    r = client.post("/ingest/file", headers=H, data={"doc_key": "REQ-TH"}, files={"file": ("req.docx", buf.getvalue())})
    assert r.status_code == 200
    assert hits(client, "radiator reject", user())[0]["doc_key"] == "REQ-TH"
    r = client.post("/ingest/file", headers=H, data={"doc_key": "PDF-1"}, files={"file": ("a.pdf", tiny_pdf("Mast deployment takes forty seconds"))})
    assert r.status_code == 200
    assert hits(client, "mast deployment", user())[0]["doc_key"] == "PDF-1"


def test_usage_is_tracked(client):
    put(client, doc_key="A", body="alpha beta gamma")
    hits(client, "alpha", user())
    ops = {u["op"] for u in client.get("/usage", headers=H).json()["usage"]}
    assert {"ingest", "search"} <= ops


def test_scanned_pdf_is_rejected_clearly(client):
    r = client.post("/ingest/file", headers=H, data={"doc_key": "SCAN"}, files={"file": ("scan.pdf", tiny_pdf(""))})
    assert r.status_code == 422 and "OCR" in r.json()["detail"]
