import io
import os

from docx import Document
from pypdf import PdfReader


def _ocr(data: bytes) -> str:
    try:
        import pypdfium2
        import pytesseract
    except ImportError:
        raise ValueError("scanned PDF: OCR is not installed (build the rag image with WITH_OCR=1)")
    return "\n\n".join(pytesseract.image_to_string(pg.render(scale=2).to_pil()) for pg in pypdfium2.PdfDocument(data))


def parse(name: str, data: bytes) -> str:
    """Extract text from PDF, DOCX (including tables, where requirements often live), or plain text."""
    n = name.lower()
    if n.endswith(".pdf"):
        text = "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        return text if text.strip() or os.getenv("OCR", "0") != "1" else _ocr(data)
    if n.endswith(".docx"):
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(x for x in parts if x.strip())
    return data.decode("utf-8", errors="replace")  # md, txt, source files
